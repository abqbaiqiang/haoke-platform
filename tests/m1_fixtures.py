"""Artificial fixtures only; never copy company exports into this module."""
import csv
import io
import openpyxl

from app.import_parser import BALANCE_NAMES, PROFIT_NAMES


def csv_bytes(rows):
    out = io.StringIO(newline='')
    csv.writer(out).writerows(rows)
    return out.getvalue().encode('utf-8-sig')


def master(kind='customer', code='001', name='测试客户', staff='甲'):
    headers = ['客户编号', '客户名称', '销售人员'] if kind == 'customer' else ['商品编号', '商品名称', '品牌']
    return csv_bytes([headers, [code, name, staff]])


def sales(order='T001', stamp='2026-08-01 10:00:00', prices=('0.10', '0.20'), customer='001', product='P001', staff='甲'):
    heads = ['单据编号', '单据日期', '客户编码', '销售金额', '销售人员', '商品编码', '数量', '金额',
             '销售单价', '最后修改时间', '优惠金额', '优惠后金额', '收款状态', '已收款', '退货状态']
    from decimal import Decimal
    total = str(sum(map(Decimal, prices)))
    rows = [heads]
    for i, price in enumerate(prices):
        rows.append([order if i == 0 else '', '2026-08-01' if i == 0 else '', customer if i == 0 else '',
                     total if i == 0 else '', staff if i == 0 else '', product, '1', price, price,
                     stamp if i == 0 else '', '0' if i == 0 else '', total if i == 0 else '',
                     '未收款' if i == 0 else '', '0' if i == 0 else '', '未退货' if i == 0 else ''])
    rows.append(['', '', '', '', '', '合计:', '', total])
    return csv_bytes(rows)


def workbook_bytes(w):
    out = io.BytesIO()
    w.save(out)
    w.close()
    return out.getvalue()


def finance(kind='profit', blank=False, revenue=100):
    w = openpyxl.Workbook()
    s = w.active
    s.title = '测试报表'
    s['A1'] = '利润表' if kind == 'profit' else '资产负债表'
    s['A2'] = '编制单位：测试公司'
    if kind == 'profit':
        s['B2'] = '2026年8期'
        s['D2'] = '单位：元'
        for c, v in enumerate(['项目', '行次', '本年累计金额', '本月金额'], 1):
            s.cell(3, c, v)
        for n in range(1, 33):
            r = n+3
            s.cell(r, 1, PROFIT_NAMES.get(n, ('', f'测试子科目{n}'))[1])
            s.cell(r, 2, str(n))
            for c in [3, 4]:
                s.cell(r, c, {1: revenue, 2: 60, 21: revenue-60, 30: revenue-60, 32: revenue-60}.get(n, 0))
        if blank:
            s['D6'] = None
    else:
        s['E2'] = '2026-08-31'
        s['F2'] = '单位：元'
        for c, v in enumerate(['资产', '行次', '期末余额', '年初余额', '负债和权益', '行次', '期末余额', '年初余额'], 1):
            s.cell(3, c, v)
        for n in range(1, 54):
            r, c = (n+3, 1) if n <= 30 else (n-27, 5)
            s.cell(r, c, BALANCE_NAMES.get(n, ('', f'测试资产科目{n}'))[1])
            s.cell(r, c+1, str(n))
            for col in [c+2, c+3]:
                s.cell(r, col, {1: 100, 15: 100, 30: 100, 31: 30, 41: 30, 47: 30, 51: 70, 52: 70, 53: 100}.get(n, 0))
    return workbook_bytes(w)
