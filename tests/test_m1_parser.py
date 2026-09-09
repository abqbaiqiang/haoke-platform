import io
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

from app.import_parser import ParseError, parse, read_book
from m1_fixtures import csv_bytes, finance, master, sales, workbook_bytes

pytestmark = pytest.mark.unit


def test_order_continuations_totals_decimal_and_repeated_sku():
    out = parse(sales(prices=('0.10', '0.10')), 'sales.csv', 'sales')
    assert not out.errors
    assert out.summary['source_amount'] == '0.20'
    assert out.summary['orders'] == 1 and out.summary['lines'] == 2
    assert out.raw[-1]['parse_status'] == 'ignored'
    assert out.data['records'][0]['customer_code'] == '001'


@pytest.mark.parametrize('body', [sales().replace(b'0.30', b'9.00'), sales().rsplit(b'\r\n', 2)[0],
                                 b'bad,header\r\nfoo,bar', sales().replace(b'T001,', b',', 1)])
def test_malformed_orders_never_silently_succeed(body):
    try:
        result = parse(body, 'bad.csv', 'sales')
        assert result.errors
    except ParseError:
        pass


def test_master_keys_and_missing_columns():
    out = parse(master(), 'customer.csv', 'customer')
    assert out.data['records'][0]['code'] == '001'
    assert parse(csv_bytes([['客户编号', '客户名称'], ['001', 'A'], ['001', 'B']]), 'x.csv', 'customer').errors
    with pytest.raises(ParseError, match='关键列'):
        parse(b'code,name\r\n1,A', 'bad.csv', 'customer')
    out = parse(b'code,name\r\n001,A', 'custom.csv', 'customer', {'_header_row': 1, 'code': 1, 'name': 2})
    assert out.data['records'][0]['code'] == '001'


def test_wrong_xlsx_dimensions_and_extra_empty_sheets():
    w = openpyxl.Workbook()
    s = w.active
    s.append(['商品编号', '商品名称'])
    s.append(['0001', '演示商品'])
    w.create_sheet('空表')
    data = workbook_bytes(w)
    out = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as z, ZipFile(out, 'w', ZIP_DEFLATED) as new:
        for name in z.namelist():
            value = z.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                value = value.replace(b'ref="A1:B2"', b'ref="A1"')
            new.writestr(name, value)
    parsed = parse(out.getvalue(), 'product.xlsx', 'product')
    assert parsed.data['records'][0]['code'] == '0001'


@pytest.mark.parametrize('data,name', [(b'html', 'fake.xls'), (b'abc', 'fake.xlsx'), (b'\x00', 'x.csv'), (b'foo', 'x.exe')])
def test_file_content_signatures(data, name):
    with pytest.raises(ParseError):
        read_book(data, name)


@pytest.mark.parametrize('kind', ['profit', 'balance_sheet'])
def test_financial_column_semantics(kind):
    out = parse(finance(kind), 'test.xlsx', kind)
    assert out.data['period'] == '2026-08-01'
    assert not out.errors
    bycode = {r['metric_code']: r for r in out.data['records']}
    if kind == 'profit':
        assert bycode['revenue']['period_value'] == '100'
        assert bycode['net_profit']['period_value'] == '40'
        assert bycode['revenue']['end_value'] is None
    else:
        assert bycode['cash']['end_value'] == '100'
        assert out.mapping['prior_basis'] == 'year_opening'
        assert bycode['cash']['period_value'] is None


def test_finance_blank_policy_and_formula_missing_cache():
    out = parse(finance(blank=True), 'test.xlsx', 'profit')
    assert out.summary['requires_blank_confirmation']
    out = parse(finance(blank=True), 'test.xlsx', 'profit', blank_as_zero=True)
    assert not out.summary['requires_blank_confirmation']
    w = openpyxl.load_workbook(io.BytesIO(finance()))
    w.active['D4'] = '=10*10'
    with pytest.raises(ParseError, match='公式'):
        parse(workbook_bytes(w), 'test.xlsx', 'profit')


def test_finance_mismatch_and_reversed_month_columns():
    w = openpyxl.load_workbook(io.BytesIO(finance()))
    w.active['D35'] = 999
    with pytest.raises(ParseError, match='勾稽不平'):
        parse(workbook_bytes(w), 'test.xlsx', 'profit')
    w = openpyxl.load_workbook(io.BytesIO(finance()))
    w.active['C3'], w.active['D3'] = '本月金额', '本年累计金额'
    with pytest.raises(ParseError, match='表头'):
        parse(workbook_bytes(w), 'test.xlsx', 'profit')


def test_malicious_xml_entity_is_rejected():
    out = io.BytesIO()
    with ZipFile(out, 'w', ZIP_DEFLATED) as z:
        z.writestr('xl/worksheets/sheet1.xml', '<!DOCTYPE foo [<!ENTITY x "bad">]><foo>&x;</foo>')
    with pytest.raises(ParseError):
        read_book(out.getvalue(), 'entity.xlsx')


def test_balance_subtotal_error_is_detected_even_when_grand_totals_balance():
    w = openpyxl.load_workbook(io.BytesIO(finance('balance_sheet')))
    w.active['C4'] = 99  # Cash component; asset and liability grand totals still both 100.
    with pytest.raises(ParseError, match='勾稽不平'):
        parse(workbook_bytes(w), 'balance.xlsx', 'balance_sheet')
