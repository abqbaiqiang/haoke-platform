"""Bounded readers and pure normalization. No filesystem writes or database access."""
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import openpyxl
import xlrd
from defusedxml import ElementTree

from app.config import get_settings

VERSION = "m1-template-1"
KINDS = {"customer", "product", "sales", "profit", "balance_sheet"}
MAX_ROWS = 50000
MAX_COLS = 256
MAX_CELLS = 1500000
# 源文件无时区信息时按展示时区（Asia/Shanghai）解释，与 docs/03 时区口径一致。
SOURCE_TZ = ZoneInfo(get_settings().app_timezone)


class ParseError(ValueError):
    def __init__(self, message, row=0, field="文件"):
        super().__init__(message)
        self.row, self.field = row, field


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def amount(value, row, name, required=True, places=2):
    if not text(value):
        if required:
            raise ParseError("必需金额/数量为空", row, name)
        return None
    try:
        result = Decimal(text(value).replace(",", ""))
        if not result.is_finite() or abs(result) >= Decimal(10) ** (18-places):
            raise InvalidOperation
        if result != result.quantize(Decimal(10) ** -places):
            raise InvalidOperation
        return str(result)
    except InvalidOperation:
        raise ParseError("数值格式、精度或范围不符合要求（公式结果不可直接作为输入）", row, name) from None


def day(value, row):
    try:
        return date.fromisoformat(text(value)[:10].replace("/", "-")).isoformat()
    except ValueError:
        raise ParseError("日期应为 YYYY-MM-DD", row, "日期") from None


def source_time(value, row):
    if not text(value):
        return None
    try:
        dt = datetime.fromisoformat(text(value).replace("/", "-"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=SOURCE_TZ)
        return dt.isoformat()
    except ValueError:
        raise ParseError("源修改时间无效", row, "最后修改时间") from None


def read_book(data: bytes, filename: str):
    suffix = PurePath(filename).suffix.lower()
    sheets = []
    try:
        if suffix == ".xlsx":
            if not data.startswith(b"PK"):
                raise ParseError("文件内容与 xlsx 扩展名不符")
            with ZipFile(io.BytesIO(data)) as z:
                infos = z.infolist()
                if len(infos) > 300 or len({x.filename for x in infos}) != len(infos):
                    raise ParseError("工作簿压缩结构异常")
                if sum(x.file_size for x in infos) > 80 * 1024 * 1024:
                    raise ParseError("工作簿展开体积超限")
                if any(x.flag_bits & 1 or 'vbaproject' in x.filename.lower() for x in infos):
                    raise ParseError("不接受加密工作簿或宏")
                count = 0
                for info in infos:
                    if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                        with z.open(info) as handle:
                            for _, node in ElementTree.iterparse(handle, events=("end",)):
                                tag = node.tag.rsplit("}", 1)[-1]
                                if tag == "row" and int(node.attrib.get("r", 0)) > MAX_ROWS:
                                    raise ParseError("工作表行号超限")
                                if tag == "c":
                                    count += 1
                                    ref = node.attrib.get("r", "")
                                    match = re.fullmatch(r"([A-Z]+)([0-9]+)", ref)
                                    if not match or openpyxl.utils.column_index_from_string(match[1]) > MAX_COLS:
                                        raise ParseError("工作表列数超限")
                                node.clear()
                if count > MAX_CELLS:
                    raise ParseError("工作簿单元格数超限")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                if len(wb.sheetnames) > 10:
                    raise ParseError("工作表数量超限")
                expanded_cells = 0
                for s in wb:
                    s.reset_dimensions()
                    rows = []
                    for row in s.iter_rows():
                        expanded_cells += len(row)
                        if len(rows) >= MAX_ROWS or len(row) > MAX_COLS or expanded_cells > MAX_CELLS:
                            raise ParseError("工作表尺寸超限")
                        rows.append([c.value for c in row])
                    sheets.append((s.title, rows))
            finally:
                wb.close()
        elif suffix == ".xls":
            if not data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
                raise ParseError("文件内容与 xls 扩展名不符")
            wb = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                count = 0
                if wb.nsheets > 10:
                    raise ParseError("工作表数量超限")
                for s in wb.sheets():
                    count += s.nrows * s.ncols
                    if s.nrows > MAX_ROWS or s.ncols > MAX_COLS or count > MAX_CELLS:
                        raise ParseError("工作表尺寸超限")
                    rows = [[xlrd.xldate_as_datetime(c.value, wb.datemode) if c.ctype == xlrd.XL_CELL_DATE
                             else c.value for c in s.row(i)] for i in range(s.nrows)]
                    sheets.append((s.name, rows))
            finally:
                wb.release_resources()
        elif suffix == ".csv":
            try:
                decoded = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ParseError("CSV 请保存为 UTF-8（可带 BOM）") from None
            if '\x00' in decoded:
                raise ParseError("CSV 包含非法空字符")
            rows = []
            count = 0
            for row in csv.reader(io.StringIO(decoded)):
                count += len(row)
                if len(rows) >= MAX_ROWS or len(row) > MAX_COLS or count > MAX_CELLS:
                    raise ParseError("CSV 尺寸超限")
                rows.append(row)
            sheets = [("CSV", rows)]
        else:
            raise ParseError("仅支持 xls、xlsx、csv")
    except ParseError:
        raise
    except (Exception,) as exc:
        # Parser exceptions may embed file contents. Return a safe, useful error only.
        raise ParseError("无法读取工作簿，请检查是否损坏、加密或格式不受支持") from exc
    if sum(len(r) for _, rows in sheets for r in rows) > MAX_CELLS:
        raise ParseError("工作簿尺寸超限")
    return sheets


@dataclass
class Parsed:
    data: dict = field(default_factory=lambda: {"records": []})
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    raw: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)
    mapping: dict = field(default_factory=dict)


ALIASES = {
    'customer': {'code': ['*客户编号', '客户编号', '客户编码'], 'name': ['*客户名称', '客户名称'], 'salesperson': ['销售人员']},
    'product': {'code': ['*商品编号', '商品编号', '商品编码'], 'name': ['商品名称']},
    'sales': {'order_no': ['单据编号'], 'order_date': ['单据日期'], 'customer_code': ['客户编码', '客户编号'],
              'salesperson': ['销售人员'], 'sales_amount': ['销售金额'], 'discount': ['优惠金额'],
              'net_amount': ['优惠后金额'], 'received': ['已收款'], 'payment_status': ['收款状态'],
              'product_code': ['商品编码', '商品编号'], 'quantity': ['数量'], 'unit': ['单位'],
              'unit_price': ['销售单价'], 'line_amount': ['金额'], 'warehouse': ['仓库'],
              'updated': ['最后修改时间'], 'return_status': ['退货状态'], 'status': ['单据状态']},
}
REQUIRED = {'customer': {'code', 'name'}, 'product': {'code', 'name'},
            'sales': {'order_no', 'order_date', 'customer_code', 'sales_amount', 'product_code', 'quantity', 'line_amount'}}


def parse(data, filename, kind, overrides=None, blank_as_zero=False):
    if kind not in KINDS:
        raise ParseError("不支持该业务类型")
    sheets = read_book(data, filename)
    result = Parsed()
    for sheet, rows in sheets:
        result.raw.extend({'sheet_name': sheet, 'row_no': i, 'raw_data': {'cells': [
            v.isoformat() if isinstance(v, (datetime, date)) else v for v in row]},
            'parse_status': 'ignored'} for i, row in enumerate(rows, 1))
    nonempty = [(s, rows) for s, rows in sheets if any(any(text(v) for v in r) for r in rows)]
    if len(nonempty) != 1:
        raise ParseError("请选择仅包含一个有效数据工作表的文件")
    sheet, rows = nonempty[0]
    if kind in {'profit', 'balance_sheet'}:
        parse_finance(result, sheet, rows, kind, blank_as_zero)
    else:
        parse_business(result, sheet, rows, kind, overrides or {})
    errors_by_row = {e['row']: e['message'] for e in result.errors}
    for row in result.raw:
        if row['sheet_name'] == sheet:
            n = row['row_no']
            if n in errors_by_row:
                row.update(parse_status='error', error_message=errors_by_row[n])
            elif n in result.summary.get('business_rows', []):
                row['parse_status'] = 'ok'
    result.summary.update(total_rows=len(result.raw), error_count=len(result.errors),
                          business_rows_count=len(result.summary.get('business_rows', [])),
                          template_version=VERSION, sheet=sheet)
    return result


def parse_business(out, sheet, rows, kind, overrides):
    aliases = ALIASES[kind]
    header = None
    for i, r in enumerate(rows[:20]):
        labels = [text(v) for v in r]
        found = {k: next((j for j, v in enumerate(labels) if v in names), None) for k, names in aliases.items()}
        if all(found[k] is not None for k in REQUIRED[kind]):
            header, mapping = i, found
            break
    if header is None:
        # Explicit mapping uses a separately declared header row; never silently guess.
        if not overrides or '_header_row' not in overrides:
            raise ParseError("缺少关键列，无法识别表头；可通过字段映射指定列")
        header = overrides['_header_row'] - 1
        mapping = {k: None for k in aliases}
    if header < 0 or header >= min(len(rows), 20):
        raise ParseError("表头行号超出范围")
    for k, v in overrides.items():
        if k == '_header_row':
            continue
        if k not in aliases or not isinstance(v, int) or v < 1 or v > len(rows[header]):
            raise ParseError("字段映射不符合要求")
        mapping[k] = v - 1
    if any(mapping[k] is None for k in REQUIRED[kind]):
        raise ParseError("字段映射缺少必需字段")
    required_cols = [mapping[k] for k in REQUIRED[kind]]
    if len(set(required_cols)) != len(required_cols):
        raise ParseError("必需字段不能映射到同一列")
    out.mapping = {'header_row': header+1, 'columns': {k: v+1 for k, v in mapping.items() if v is not None}}
    out.summary['business_rows'] = []
    records, seen, current = out.data['records'], set(), None

    def get(r, k):
        j = mapping.get(k)
        return r[j] if j is not None and j < len(r) else None

    def required(r, k, n):
        value = text(get(r, k))
        if not value or len(value) > (255 if k == 'name' else 100) or value.startswith('='):
            raise ParseError("必需文本字段为空、过长或为公式", n, k)
        return value

    def issue(e):
        out.errors.append({'sheet': sheet, 'row': e.row, 'field': e.field, 'message': str(e)})

    def close_order(n):
        nonlocal current
        if current:
            issue(ParseError("订单缺少合计行，无法确认明细完整性", current['row'], '合计'))
        current = None

    for index, r in enumerate(rows[header+1:], header+2):
        try:
            if not any(text(v) for v in r):
                if kind == 'sales':
                    close_order(index)
                continue
            if kind != 'sales':
                code = required(r, 'code', index)
                if code in seen:
                    raise ParseError("文件内编码重复，请核对主档", index, 'code')
                seen.add(code)
                records.append({'row': index, 'code': code, 'name': required(r, 'name', index),
                                'salesperson': text(get(r, 'salesperson'))})
            else:
                product_code = text(get(r, 'product_code'))
                if product_code.rstrip(':：') in {'合计', '小计', '总计'} and not text(get(r, 'order_no')):
                    if not current:
                        raise ParseError("合计行没有对应订单", index, '合计')
                    subtotal = Decimal(amount(get(r, 'line_amount'), index, '合计金额'))
                    line_sum = sum((Decimal(x['line_amount']) for x in current['lines']), Decimal(0))
                    if not current['lines'] or line_sum != subtotal or line_sum != Decimal(current['sales_amount']):
                        raise ParseError("订单头、明细和合计金额不一致", index, '金额勾稽')
                    current = None
                    continue
                if text(get(r, 'order_no')):
                    close_order(index)
                    key = required(r, 'order_no', index)
                    if key in seen:
                        raise ParseError("文件内订单号重复", index, '单据编号')
                    seen.add(key)
                    if text(get(r, 'return_status')) not in {'', '未退货'}:
                        raise ParseError("退货状态需先确认映射，不能按普通销售静默导入", index, '退货状态')
                    if text(get(r, 'status')) not in {'', '正常', '有效'}:
                        raise ParseError("单据状态尚未配置，需核对作废/取消规则", index, '单据状态')
                    current = {'row': index, 'order_no': key, 'order_date': day(get(r, 'order_date'), index),
                               'customer_code': required(r, 'customer_code', index), 'salesperson': text(get(r, 'salesperson')),
                               'sales_amount': amount(get(r, 'sales_amount'), index, '销售金额'),
                               'discount_amount': amount(get(r, 'discount'), index, '优惠金额', False),
                               'net_amount': amount(get(r, 'net_amount'), index, '优惠后金额', False),
                               'received_amount': amount(get(r, 'received'), index, '已收款', False),
                               'payment_status': text(get(r, 'payment_status')), 'source_updated_at': source_time(get(r, 'updated'), index),
                               'lines': []}
                    if current['net_amount'] is not None and current['discount_amount'] is not None:
                        if Decimal(current['sales_amount'])-Decimal(current['discount_amount']) != Decimal(current['net_amount']):
                            raise ParseError("销售金额减优惠与优惠后金额不一致", index, '优惠')
                    records.append(current)
                if current is None:
                    raise ParseError("孤立明细或未知行，请检查订单边界", index, '单据编号')
                line = {'row': index, 'product_code': required(r, 'product_code', index),
                        'quantity': amount(get(r, 'quantity'), index, '数量', places=4),
                        'unit_price': amount(get(r, 'unit_price'), index, '单价', False, 4),
                        'line_amount': amount(get(r, 'line_amount'), index, '明细金额'),
                        'unit_name': text(get(r, 'unit'))[:100] or None,
                        'warehouse_name': text(get(r, 'warehouse'))[:255] or None}
                if Decimal(line['quantity']) <= 0 or Decimal(line['line_amount']) < 0:
                    raise ParseError("非正数量或负金额需退货规则，不按普通销售导入", index, '数量/金额')
                current['lines'].append(line)
            out.summary['business_rows'].append(index)
        except ParseError as exc:
            issue(exc)
            if kind == 'sales':
                current = None
    if kind == 'sales':
        close_order(len(rows))
        dates = [r['order_date'] for r in records]
        out.summary.update(orders=len(records), lines=sum(len(r['lines']) for r in records),
                           source_amount=str(sum((Decimal(r['sales_amount']) for r in records), Decimal(0))),
                           date_from=min(dates) if dates else None, date_to=max(dates) if dates else None)
        out.warnings.append('销售金额为源文件核对值；退货、作废、税价及完整性未验证，不等同最终净经营销售额。')
        out.warnings.append('本次欠款仅保留原始值，不作为当前应收；当前采购价不作为历史成本。')
        zeros = sum(Decimal(x['line_amount']) == 0 for r in records for x in r['lines'])
        if zeros:
            out.warnings.append(f'{zeros} 条零金额明细原样保留，请复核是否赠品或业务调价。')
    out.summary['records'] = len(records)
    if not records:
        issue(ParseError("没有可导入的业务记录"))


PROFIT_NAMES = {1: ('revenue', '营业收入'), 2: ('cost', '营业成本'), 3: ('tax_surcharge', '税金及附加'),
                11: ('selling_expense', '销售费用'), 14: ('management_expense', '管理费用'),
                18: ('finance_expense', '财务费用'), 20: ('investment_income', '投资收益'),
                21: ('operating_profit', '营业利润'), 22: ('nonoperating_income', '营业外收入'),
                24: ('nonoperating_expense', '营业外支出'), 30: ('profit_total', '利润总额'),
                31: ('income_tax', '所得税费用'), 32: ('net_profit', '净利润')}
BALANCE_NAMES = {1: ('cash', '货币资金'), 4: ('ar', '应收账款'), 9: ('inventory', '存货'),
                 15: ('current_assets', '流动资产合计'), 29: ('noncurrent_assets', '非流动资产合计'),
                 30: ('assets', '资产总计'), 41: ('current_liabilities', '流动负债合计'),
                 46: ('noncurrent_liabilities', '非流动负债合计'), 47: ('liabilities', '负债合计'),
                 52: ('equity', '所有者权益'), 53: ('liabilities_equity', '总计')}


def parse_finance(out, sheet, rows, kind, blank_as_zero):
    if len(rows) < 5:
        raise ParseError("财务报表缺少标题和数据")
    def cell(r, c):
        return rows[r-1][c-1] if r <= len(rows) and c <= len(rows[r-1]) else None
    title = text(cell(1, 1))
    if title != ('利润表' if kind == 'profit' else '资产负债表'):
        raise ParseError("报表类型与所选模板不符")
    entity = text(cell(2, 1)).removeprefix('编制单位：').removeprefix('编制单位:').strip()
    if not entity:
        raise ParseError("缺少编制单位", 2, '编制单位')
    if kind == 'profit':
        match = re.fullmatch(r'(\d{4})年(\d{1,2})期', text(cell(2, 2)))
        if not match:
            raise ParseError("利润表期间无法识别", 2, '期间')
        period = day(f'{match[1]}-{int(match[2]):02d}-01', 2)
        if text(cell(2, 4)) != '单位：元' or [text(cell(3, c)) for c in [3, 4]] != ['本年累计金额', '本月金额']:
            raise ParseError("单位或本月/累计表头与模板不符", 3, '表头')
        specs = [(1, 2, 4, 3)]
        names = PROFIT_NAMES
    else:
        period = day(cell(2, 5), 2)[:8] + '01'
        if text(cell(2, 6)) != '单位：元' or [text(cell(3, c)) for c in [3, 4, 7, 8]] != ['期末余额', '年初余额']*2:
            raise ParseError("单位或期末/年初表头与模板不符", 3, '表头')
        specs = [(1, 2, 3, 4), (5, 6, 7, 8)]
        names = BALANCE_NAMES
    seen = set()
    out.summary['business_rows'] = []
    vals = {}
    blanks = 0
    for r in range(4, len(rows)+1):
        for label_col, no_col, current_col, prior_col in specs:
            if not text(cell(r, no_col)):
                continue
            try:
                no = int(text(cell(r, no_col)))
            except ValueError:
                raise ParseError("行次无效", r, '行次') from None
            if no in seen or no < 1 or no > (32 if kind == 'profit' else 53):
                raise ParseError("科目行次重复或不受支持", r, '行次')
            seen.add(no)
            name = text(cell(r, label_col))
            if not name or (no in names and names[no][1] not in name):
                raise ParseError("科目名称与行次不匹配", r, '科目')
            code = names[no][0] if no in names else f'{kind}_row_{no}'
            current = amount(cell(r, current_col), r, name, False)
            prior = amount(cell(r, prior_col), r, name, False)
            blanks += int(current is None)+int(prior is None)
            vals[no] = (current, prior)
            rec = {'row': r, 'metric_code': code, 'metric_name': name, 'period_value': None, 'ytd_value': None,
                   'begin_value': None, 'end_value': None}
            if kind == 'profit':
                rec.update(period_value=current, ytd_value=prior)
            else:
                rec.update(end_value=current, begin_value=prior)
            out.data['records'].append(rec)
            out.summary['business_rows'].append(r)
    if seen != set(range(1, 33 if kind == 'profit' else 54)):
        raise ParseError("财务模板缺少科目行次，请使用完整报表")
    required = [1, 2, 21, 30, 32] if kind == 'profit' else [1, 4, 9, 15, 30, 41, 47, 52, 53]
    if any(vals[k][0] is None for k in required):
        raise ParseError("关键财务指标为空，请填写真实值或明确的零，不能静默补零")
    checks = ([(21, [(1, 1), (2, -1), (3, -1), (11, -1), (14, -1), (18, -1), (20, 1)]),
               (30, [(21, 1), (22, 1), (24, -1)]), (32, [(30, 1), (31, -1)])] if kind == 'profit'
              else [(30, [(53, 1)]), (30, [(15, 1), (29, 1)]), (53, [(47, 1), (52, 1)]),
                    (47, [(41, 1), (46, 1)]),
                    (15, [(k, 1) for k in list(range(1, 10))+[14]]),
                    (29, [(k, 1) for k in [16, 17]+list(range(20, 29))]),
                    (41, [(k, 1) for k in range(31, 41)]),
                    (46, [(k, 1) for k in range(42, 46)]),
                    (52, [(k, 1) for k in range(48, 52)])])
    for target, terms in checks:
        for col in [0, 1]:
            if vals[target][col] is None:
                continue
            if any(vals[k][col] is None for k, _ in terms) and not blank_as_zero:
                out.warnings.append('存在空白科目：需确认“空白按零参与勾稽”后才能完成财务确认；原值仍保留空。')
                continue
            expected = sum((Decimal(vals[k][col] or '0')*sign for k, sign in terms), Decimal(0))
            if expected != Decimal(vals[target][col]):
                raise ParseError("财务金额勾稽不平，不能正式导入", 0, names[target][1])
    if blanks:
        out.warnings.append(f'{blanks} 个金额单元格为空，空白与零分开保存。')
    out.warnings = list(dict.fromkeys(out.warnings))
    out.data.update(period=period, entity_name=entity)
    out.summary.update(records=len(seen), date_from=period, date_to=period, blank_count=blanks,
                       requires_blank_confirmation=bool(blanks and not blank_as_zero),
                       business_rows=sorted(set(out.summary['business_rows'])))
    out.mapping = {'version': VERSION, 'type': kind, 'current_columns': [4] if kind == 'profit' else [3, 7],
                   'prior_columns': [3] if kind == 'profit' else [4, 8], 'prior_basis': 'ytd' if kind == 'profit' else 'year_opening'}
